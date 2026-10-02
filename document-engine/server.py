import os
import shutil
import subprocess
import tempfile
import uuid
from datetime import datetime
from pathlib import Path

from fastapi import FastAPI, Header, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.background import BackgroundTasks
from pydantic import BaseModel
from docx import Document
from supabase import create_client, Client


# =========================================================
# CONFIG
# =========================================================

SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_SERVICE_ROLE_KEY = os.getenv("SUPABASE_SERVICE_ROLE_KEY")
ENGINE_API_KEY = os.getenv("ENGINE_API_KEY")

TEMPLATE_BUCKET = "templates"
TEMPLATE_PATH = "SPKL_Arif Chandra F.docx"


if not SUPABASE_URL:
    raise RuntimeError("SUPABASE_URL belum diatur")

if not SUPABASE_SERVICE_ROLE_KEY:
    raise RuntimeError("SUPABASE_SERVICE_ROLE_KEY belum diatur")

if not ENGINE_API_KEY:
    raise RuntimeError("ENGINE_API_KEY belum diatur")


supabase: Client = create_client(
    SUPABASE_URL,
    SUPABASE_SERVICE_ROLE_KEY,
)


# =========================================================
# APP
# =========================================================

app = FastAPI(
    title="JOSJIS Document Engine",
    version="1.0.0",
)


app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["GET", "POST", "OPTIONS"],
    allow_headers=["*"],
)


# =========================================================
# MODELS
# =========================================================

class GenerateSpklRequest(BaseModel):
    lembur_ids: list[str]


# =========================================================
# CONSTANTS
# =========================================================

MONTHS = [
    "Januari",
    "Februari",
    "Maret",
    "April",
    "Mei",
    "Juni",
    "Juli",
    "Agustus",
    "September",
    "Oktober",
    "November",
    "Desember",
]

WEEKDAYS = [
    "Senin",
    "Selasa",
    "Rabu",
    "Kamis",
    "Jumat",
    "Sabtu",
    "Minggu",
]

# Nilai otomatis untuk kolom:
# CATATAN PENGAWASAN*)
SUPERVISION_NOTE = "Dilaksanakan dengan baik"


# =========================================================
# UNIT MAPPING
# =========================================================

UNIT_NAME_BY_ID = {
    "e2b50285-3bd4-4f78-bbde-c299f92ae1e8": "KCP KUAMANG KUNING 1",
    "0baec32b-1845-45f3-bcc0-d2202b0b5976": "KCP KUAMANG KUNING 2",
}


# =========================================================
# DATE / TIME HELPERS
# =========================================================

def parse_date(value):
    if not value:
        return None

    text = str(value)

    if len(text) >= 10:
        try:
            return datetime.fromisoformat(text[:10])
        except ValueError:
            pass

    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return None


def format_date_long(value):
    date = parse_date(value)

    if not date:
        return ""

    return (
        f"{date.day} "
        f"{MONTHS[date.month - 1]} "
        f"{date.year}"
    )


def format_weekday(value):
    date = parse_date(value)

    if not date:
        return ""

    return WEEKDAYS[date.weekday()]


def format_time(value):
    if not value:
        return ""

    text = str(value)

    return text[:5]


# =========================================================
# TEMPLATE
# =========================================================

def download_template(destination: Path):
    try:
        response = (
            supabase.storage
            .from_(TEMPLATE_BUCKET)
            .download(TEMPLATE_PATH)
        )
    except Exception as exc:
        raise RuntimeError(
            f"Gagal mengambil template dari Supabase Storage: {exc}"
        )

    destination.write_bytes(response)


# =========================================================
# DATABASE
# =========================================================

def get_lembur_records(lembur_ids):
    result = (
        supabase
        .table("lembur_mka")
        .select(
            """
            id,
            unit_id,
            mka_id,
            pengawas_id,
            tanggal,
            jam_mulai,
            jam_selesai,
            pekerjaan,

            pegawai_mka:pegawai!lembur_mka_mka_id_fkey(
                id,
                nip,
                nama,
                jabatan,
                jenis_pegawai,
                unit_id
            ),

            pegawai_pengawas:pegawai!lembur_mka_pengawas_id_fkey(
                id,
                nip,
                nama,
                jabatan,
                jenis_pegawai,
                unit_id
            )
            """
        )
        .in_("id", lembur_ids)
        .execute()
    )

    records = result.data or []

    if not records:
        raise HTTPException(
            status_code=404,
            detail="Data lembur tidak ditemukan.",
        )

    # Pertahankan urutan sesuai lembur_ids
    indexed = {
        str(item["id"]): item
        for item in records
    }

    ordered = []

    for lembur_id in lembur_ids:
        item = indexed.get(str(lembur_id))

        if item:
            ordered.append(item)

    if len(ordered) != len(lembur_ids):
        raise HTTPException(
            status_code=404,
            detail="Sebagian data lembur tidak ditemukan.",
        )

    return ordered


def get_branch_manager():
    """
    Mengambil Branch Manager dari Master Pegawai.
    """

    try:
        result = (
            supabase
            .table("pegawai")
            .select("*")
            .eq("jenis_pegawai", "BRANCH_MANAGER")
            .execute()
        )

        employees = result.data or []

    except Exception as exc:
        raise RuntimeError(
            f"Gagal mengambil data Branch Manager: {exc}"
        )

    if not employees:
        return {}

    # Cari yang aktif.
    active_candidates = []

    for employee in employees:
        status_found = False
        is_active = True

        for key in [
            "aktif",
            "is_active",
            "active",
            "status_aktif",
        ]:
            if key in employee:
                status_found = True
                value = employee.get(key)

                if isinstance(value, bool):
                    is_active = value

                elif isinstance(value, str):
                    is_active = (
                        value.strip().lower()
                        in [
                            "aktif",
                            "active",
                            "true",
                            "1",
                            "yes",
                        ]
                    )

                elif isinstance(value, int):
                    is_active = value == 1

                break

        # Kalau database tidak memiliki kolom status,
        # anggap record aktif.
        if not status_found:
            is_active = True

        if is_active:
            active_candidates.append(employee)

    if not active_candidates:
        return {}

    # Prioritaskan BM yang unit_id NULL
    # karena Branch Manager berada di level cabang.
    for employee in active_candidates:
        if not employee.get("unit_id"):
            return employee

    return active_candidates[0]


# =========================================================
# WORD TEXT HELPERS
# =========================================================

def replace_in_paragraph(paragraph, replacements):
    full_text = paragraph.text

    if not full_text:
        return

    changed = full_text

    for old, new in replacements.items():
        changed = changed.replace(
            old,
            new,
        )

    if changed == full_text:
        return

    paragraph.text = changed


def set_cell_text(cell, value):
    cell.text = "" if value is None else str(value)


# =========================================================
# TABLE DETECTION
# =========================================================

# Berdasarkan pemeriksaan template asli:
#
# TABLE 0 = Petugas Kerja Lembur
# TABLE 1 = Pengawas Kerja Lembur
# TABLE 2 = Menyetujui / Pengawas / Atasan Langsung
#
# Struktur:
#
# TABLE 0 = 8 kolom
# TABLE 1 = 7 kolom
# TABLE 2 = 2 kolom
#
# Tidak perlu lagi mencari header secara tekstual.
# =========================================================

def get_worker_table(document):
    if len(document.tables) < 1:
        raise RuntimeError(
            "Tabel Petugas Kerja Lembur tidak ditemukan."
        )

    return document.tables[0]


def get_supervisor_table(document):
    if len(document.tables) < 2:
        raise RuntimeError(
            "Tabel Pengawas Kerja Lembur tidak ditemukan."
        )

    return document.tables[1]


def get_approval_table(document):
    if len(document.tables) < 3:
        return None

    return document.tables[2]


# =========================================================
# UNIT
# =========================================================

def get_work_unit(record):
    """
    Menghasilkan nama unit kerja.
    """

    unit_id = record.get("unit_id")

    if not unit_id:
        worker = record.get("pegawai_mka") or {}
        unit_id = worker.get("unit_id")

    if not unit_id:
        return ""

    unit_name = UNIT_NAME_BY_ID.get(
        str(unit_id)
    )

    if unit_name:
        return unit_name

    # Jangan tampilkan UUID di dokumen.
    return ""


# =========================================================
# WORKER TABLE
# =========================================================

def fill_worker_table(document, record):
    table = get_worker_table(document)

    if len(table.rows) < 2:
        raise RuntimeError(
            "Tabel Petugas Kerja Lembur tidak memiliki baris data."
        )

    if len(table.columns) < 8:
        raise RuntimeError(
            "Struktur tabel Petugas Kerja Lembur tidak sesuai template."
        )

    worker = record.get("pegawai_mka") or {}

    row = table.rows[1]

    values = [
        "1",
        worker.get("nip", ""),
        worker.get("nama", ""),
        worker.get("jabatan", ""),
        get_work_unit(record),
        record.get("pekerjaan", ""),
        (
            f"{format_time(record.get('jam_mulai'))}"
            f"/"
            f"{format_time(record.get('jam_selesai'))}"
        ),
        "",
    ]

    for index, value in enumerate(values):
        if index < len(row.cells):
            set_cell_text(
                row.cells[index],
                value,
            )


# =========================================================
# SUPERVISOR TABLE
# =========================================================

def fill_supervisor_table(document, record):
    table = get_supervisor_table(document)

    if len(table.rows) < 2:
        raise RuntimeError(
            "Tabel Pengawas Kerja Lembur tidak memiliki baris data."
        )

    if len(table.columns) < 7:
        raise RuntimeError(
            "Struktur tabel Pengawas Kerja Lembur tidak sesuai template."
        )

    supervisor = (
        record.get("pegawai_pengawas")
        or {}
    )

    row = table.rows[1]

    values = [
        # Kolom 0 - No.
        "1",

        # Kolom 1 - NIP
        supervisor.get("nip", ""),

        # Kolom 2 - NAMA
        supervisor.get("nama", ""),

        # Kolom 3 - JABATAN
        supervisor.get("jabatan", ""),

        # Kolom 4 - UNIT KERJA
        get_work_unit(record),

        # Kolom 5 - TANDA TANGAN
        "",

        # Kolom 6 - CATATAN PENGAWASAN*)
        SUPERVISION_NOTE,
    ]

    for index, value in enumerate(values):
        if index < len(row.cells):
            set_cell_text(
                row.cells[index],
                value,
            )


# =========================================================
# BRANCH MANAGER / APPROVAL
# =========================================================

def fill_branch_manager(document, record):
    branch_manager = get_branch_manager()

    if not branch_manager:
        raise RuntimeError(
            "Branch Manager aktif tidak ditemukan "
            "di Master Pegawai."
        )

    nama = branch_manager.get("nama", "")

    jabatan = (
        branch_manager.get("jabatan", "")
        or "BM"
    )

    # Cari seluruh paragraph di dokumen.
    paragraphs = list(document.paragraphs)

    # Tambahkan paragraph dalam semua tabel.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(
                    cell.paragraphs
                )

    replaced_name = False
    replaced_position = False

    for paragraph in paragraphs:
        text = paragraph.text

        # Nama contoh pada template.
        if "Adil Ariyanto" in text:
            paragraph.text = text.replace(
                "Adil Ariyanto",
                nama,
            )

            replaced_name = True

        # Jabatan BM pada template.
        if text.strip() == "BM":
            paragraph.text = jabatan
            replaced_position = True

    # Tidak fatal kalau salah satu tidak ditemukan.
    # Template sudah diketahui memiliki keduanya.
    return (
        replaced_name,
        replaced_position,
    )


# =========================================================
# DOCUMENT
# =========================================================

def fill_document(document, record):
    tanggal = record.get("tanggal")

    hari = format_weekday(tanggal)

    tanggal_long = format_date_long(tanggal)

    replacements = {
        "Selasa": hari,
        "04 November 2025": tanggal_long,
        "Non Rutin": "Non Rutin",
    }

    # -----------------------------------------------------
    # Paragraph biasa
    # -----------------------------------------------------

    for paragraph in document.paragraphs:
        replace_in_paragraph(
            paragraph,
            replacements,
        )

    # -----------------------------------------------------
    # Paragraph dalam tabel
    # -----------------------------------------------------

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(
                        paragraph,
                        replacements,
                    )

    # -----------------------------------------------------
    # Isi tabel petugas
    # -----------------------------------------------------

    fill_worker_table(
        document,
        record,
    )

    # -----------------------------------------------------
    # Isi tabel pengawas
    # -----------------------------------------------------

    fill_supervisor_table(
        document,
        record,
    )

    # -----------------------------------------------------
    # Isi Branch Manager
    # -----------------------------------------------------

    fill_branch_manager(
        document,
        record,
    )


# =========================================================
# DOCX -> PDF
# =========================================================

def convert_docx_to_pdf(
    docx_path: Path,
    output_dir: Path,
):
    command = [
        "libreoffice",
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(output_dir),
        str(docx_path),
    ]

    process = subprocess.run(
        command,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        timeout=120,
    )

    if process.returncode != 0:
        raise RuntimeError(
            "Gagal convert DOCX ke PDF.\n"
            f"STDOUT: {process.stdout}\n"
            f"STDERR: {process.stderr}"
        )

    pdf_path = (
        output_dir /
        f"{docx_path.stem}.pdf"
    )

    if not pdf_path.exists():
        raise RuntimeError(
            "LibreOffice selesai tetapi PDF tidak ditemukan."
        )

    return pdf_path


# =========================================================
# MERGE PDF
# =========================================================

def merge_pdfs(
    pdf_paths,
    output_path,
):
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()

    for pdf_path in pdf_paths:
        reader = PdfReader(
            str(pdf_path)
        )

        for page in reader.pages:
            writer.add_page(page)

    with open(
        output_path,
        "wb",
    ) as output:
        writer.write(output)


# =========================================================
# CLEANUP
# =========================================================

def cleanup_work_dir(work_dir: Path):
    """
    Menghapus temporary directory setelah
    FileResponse selesai dikirim.
    """

    try:
        if work_dir.exists():
            shutil.rmtree(
                work_dir,
                ignore_errors=True,
            )
    except Exception:
        # Cleanup tidak boleh membuat response PDF gagal.
        pass


# =========================================================
# ROUTES
# =========================================================

@app.get("/")
def root():
    return {
        "service": "JOSJIS Document Engine",
        "status": "online",
    }


@app.get("/health")
def health():
    return {
        "status": "ok",
    }


@app.post("/generate/spkl")
def generate_spkl(
    payload: GenerateSpklRequest,
    background_tasks: BackgroundTasks,
    x_engine_key: str | None = Header(default=None),
):
    # =====================================================
    # API KEY
    # =====================================================

    if x_engine_key != ENGINE_API_KEY:
        raise HTTPException(
            status_code=401,
            detail="Unauthorized.",
        )

    # =====================================================
    # VALIDASI
    # =====================================================

    if not payload.lembur_ids:
        raise HTTPException(
            status_code=400,
            detail="lembur_ids wajib diisi.",
        )

    # Hilangkan ID kosong dan duplikat
    # tanpa mengubah urutan.
    cleaned_ids = []

    for lembur_id in payload.lembur_ids:
        value = str(lembur_id).strip()

        if value and value not in cleaned_ids:
            cleaned_ids.append(value)

    if not cleaned_ids:
        raise HTTPException(
            status_code=400,
            detail="lembur_ids wajib diisi.",
        )

    # =====================================================
    # WORK DIRECTORY
    # =====================================================

    work_dir = Path(
        tempfile.mkdtemp(
            prefix="josjis-spkl-"
        )
    )

    try:
        # =================================================
        # AMBIL DATA LEMBUR
        # =================================================

        records = get_lembur_records(
            cleaned_ids
        )

        pdf_paths = []

        # =================================================
        # SATU RECORD = SATU DOCX = SATU PDF
        # =================================================

        for index, record in enumerate(
            records,
            start=1,
        ):
            docx_path = (
                work_dir /
                f"SPKL-{index}.docx"
            )

            pdf_dir = (
                work_dir /
                f"pdf-{index}"
            )

            pdf_dir.mkdir(
                parents=True,
                exist_ok=True,
            )

            # -------------------------------------------------
            # Ambil template fresh untuk setiap record.
            # -------------------------------------------------

            download_template(
                docx_path
            )

            # -------------------------------------------------
            # Buka Word template.
            # -------------------------------------------------

            document = Document(
                str(docx_path)
            )

            # -------------------------------------------------
            # Isi data.
            # -------------------------------------------------

            fill_document(
                document,
                record,
            )

            # -------------------------------------------------
            # Simpan DOCX.
            # -------------------------------------------------

            document.save(
                str(docx_path)
            )

            # -------------------------------------------------
            # Convert ke PDF.
            # -------------------------------------------------

            pdf_path = convert_docx_to_pdf(
                docx_path,
                pdf_dir,
            )

            pdf_paths.append(
                pdf_path
            )

        # =================================================
        # GABUNG SEMUA PDF
        # =================================================

        output_pdf = (
            work_dir /
            f"SPKL-{uuid.uuid4().hex}.pdf"
        )

        merge_pdfs(
            pdf_paths,
            output_pdf,
        )

        # Pastikan file hasil benar-benar ada.
        if not output_pdf.exists():
            raise RuntimeError(
                "PDF SPKL berhasil diproses tetapi file hasil tidak ditemukan."
            )

        # =================================================
        # CLEANUP SETELAH RESPONSE SELESAI
        # =================================================

        background_tasks.add_task(
            cleanup_work_dir,
            work_dir,
        )

        # =================================================
        # RETURN PDF
        # =================================================

        return FileResponse(
            path=str(output_pdf),
            media_type="application/pdf",
            filename="SPKL.pdf",
        )

    except HTTPException:
        cleanup_work_dir(work_dir)
        raise

    except Exception as exc:
        cleanup_work_dir(work_dir)

        raise HTTPException(
            status_code=500,
            detail=str(exc),
        )