from docx import Document

path = "/templates/SPKL_Arif Chandra F.docx"
doc = Document(path)

print("TABLES:", len(doc.tables))

for i, table in enumerate(doc.tables):
    print()
    print("=" * 80)
    print(f"TABLE {i} | ROWS={len(table.rows)} | COLS={len(table.columns)}")
    print("=" * 80)

    for r, row in enumerate(table.rows[:5]):
        cells = []
        for cell in row.cells:
            text = cell.text.replace("\n", " | ")
            cells.append(repr(text))
        print(f"ROW {r}: {cells}")
